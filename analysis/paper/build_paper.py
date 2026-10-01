"""Render paper/content.py to Word (.docx) and PDF."""

import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # the analysis/ folder
import re, sys, html

sys.path.insert(0, ROOT + "/paper")
from content import DOC, TITLE, B

OUT = f"{B}/paper/Sloan_Paper_Draft"
MK = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def segs(t):
    for part in MK.split(t):
        if not part:
            continue
        if part.startswith("**"):
            yield part[2:-2], True, False
        elif part.startswith("*") and len(part) > 1:
            yield part[1:-1], False, True
        else:
            yield part, False, False


# ---------------- DOCX
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

d = Document()
st = d.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
for sec in d.sections:
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)


def runs(par, t, size=None, italic_all=False):
    for txt, b, i in segs(t):
        r = par.add_run(txt)
        r.bold = b
        r.italic = i or italic_all
        if size:
            r.font.size = Pt(size)


fig_n = 0
for blk in DOC:
    k = blk[0]
    if k == "title":
        p = d.add_paragraph()
        r = p.add_run(blk[1])
        r.bold = True
        r.font.size = Pt(18)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif k == "meta":
        p = d.add_paragraph()
        r = p.add_run(blk[1])
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif k == "h1":
        d.add_heading(blk[1], level=1)
    elif k == "h2":
        d.add_heading(blk[1], level=2)
    elif k == "p":
        p = d.add_paragraph()
        runs(p, blk[1])
        p.paragraph_format.space_after = Pt(6)
    elif k == "bullets":
        for b_ in blk[1]:
            p = d.add_paragraph(style="List Bullet")
            runs(p, b_)
    elif k == "eq":
        p = d.add_paragraph()
        r = p.add_run(blk[1])
        r.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif k == "fig":
        d.add_picture(blk[1], width=Inches(blk[2]))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = d.add_paragraph()
        runs(p, blk[3], size=9, italic_all=True)
    elif k == "table":
        rows = blk[1]
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, c in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                par = cell.paragraphs[0]
                r = par.add_run(c)
                r.font.size = Pt(8.5)
                r.bold = i == 0
        p = d.add_paragraph()
        runs(p, blk[2], size=9, italic_all=True)
    elif k == "refs":
        for ref in blk[1]:
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            r = p.add_run(ref)
            r.font.size = Pt(9)
    elif k == "pagebreak":
        d.add_page_break()
d.save(OUT + ".docx")
# ---------------- PDF
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (
    BaseDocTemplate,
    PageTemplate,
    Frame,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image,
    PageBreak,
    KeepTogether,
    ListFlowable,
    ListItem,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.pdfmetrics import registerFontFamily
import matplotlib

FD = matplotlib.get_data_path() + "/fonts/ttf/"
for n, f in (
    ("DV", "DejaVuSans.ttf"),
    ("DVB", "DejaVuSans-Bold.ttf"),
    ("DVI", "DejaVuSans-Oblique.ttf"),
    ("DVBI", "DejaVuSans-BoldOblique.ttf"),
):
    pdfmetrics.registerFont(TTFont(n, FD + f))
registerFontFamily("DV", normal="DV", bold="DVB", italic="DVI", boldItalic="DVBI")
INK, INK2, RULE, TINT = (colors.HexColor(c) for c in ("#0b0b0b", "#52514e", "#d9d8d3", "#f0efec"))
S = {
    "title": ParagraphStyle(
        "t", fontName="DVB", fontSize=17, leading=21, alignment=TA_CENTER, spaceAfter=6, textColor=INK
    ),
    "meta": ParagraphStyle("m", fontName="DV", fontSize=8.5, leading=11, alignment=TA_CENTER, textColor=INK2),
    "h1": ParagraphStyle(
        "h1", fontName="DVB", fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=5, textColor=INK, keepWithNext=1
    ),
    "h2": ParagraphStyle(
        "h2", fontName="DVB", fontSize=10.5, leading=13.5, spaceBefore=8, spaceAfter=3, textColor=INK, keepWithNext=1
    ),
    "p": ParagraphStyle(
        "p", fontName="DV", fontSize=9.3, leading=13.2, alignment=TA_JUSTIFY, spaceAfter=5, textColor=INK
    ),
    "eq": ParagraphStyle(
        "eq", fontName="DVI", fontSize=9.3, leading=13, alignment=TA_CENTER, spaceBefore=3, spaceAfter=6, textColor=INK
    ),
    "cap": ParagraphStyle("cap", fontName="DVI", fontSize=8, leading=10.5, textColor=INK2, spaceAfter=8),
    "cell": ParagraphStyle("c", fontName="DV", fontSize=7.6, leading=9.6, textColor=INK),
    "cellb": ParagraphStyle("cb", fontName="DVB", fontSize=7.6, leading=9.6, textColor=INK),
    "ref": ParagraphStyle(
        "r", fontName="DV", fontSize=8.2, leading=10.8, leftIndent=14, firstLineIndent=-14, spaceAfter=3, textColor=INK
    ),
}


def rl(t):
    out = ""
    for txt, b, i in segs(t):
        e = html.escape(txt, quote=False)
        out += f"<b>{e}</b>" if b else (f"<i>{e}</i>" if i else e)
    return out


def img(path, w):
    im = Image(path)
    r = im.imageHeight / im.imageWidth
    im.drawWidth = w * inch
    im.drawHeight = w * inch * r
    im.hAlign = "CENTER"
    return im


def footer(c, doc_):
    c.saveState()
    c.setFont("DV", 7)
    c.setFillColor(INK2)
    c.drawString(0.8 * inch, 0.5 * inch, TITLE + " (draft)")
    c.drawRightString(7.7 * inch, 0.5 * inch, str(doc_.page))
    c.restoreState()


doc = BaseDocTemplate(
    OUT + ".pdf",
    pagesize=letter,
    leftMargin=0.8 * inch,
    rightMargin=0.8 * inch,
    topMargin=0.7 * inch,
    bottomMargin=0.8 * inch,
    title=TITLE,
)
doc.addPageTemplates(
    [PageTemplate(frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height)], onPage=footer)]
)
F = []
W = doc.width / inch
for blk in DOC:
    k = blk[0]
    if k in ("title", "meta", "h1", "h2", "p", "eq"):
        F.append(Paragraph(rl(blk[1]), S[k]))
    elif k == "bullets":
        F.append(
            ListFlowable(
                [ListItem(Paragraph(rl(b_), S["p"]), leftIndent=12) for b_ in blk[1]],
                bulletType="bullet",
                start="•",
                leftIndent=12,
                bulletFontName="DV",
            )
        )
    elif k == "fig":
        F.append(KeepTogether([img(blk[1], min(blk[2], W)), Spacer(1, 3), Paragraph(rl(blk[3]), S["cap"])]))
    elif k == "table":
        rows, cap, cw = blk[1], blk[2], blk[3]
        data = [[Paragraph(rl(c), S["cellb"] if i == 0 else S["cell"]) for c in row] for i, row in enumerate(rows)]
        tb = Table(data, colWidths=[c * inch for c in cw] if cw else None, repeatRows=1, hAlign="CENTER")
        tb.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), TINT),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, INK2),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.3, RULE),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        if len(rows) > 14:  # long tables (claims ledger) may split across pages; the header row repeats
            F.extend([tb, Spacer(1, 3), Paragraph(rl(cap), S["cap"])])
        else:
            F.append(KeepTogether([tb, Spacer(1, 3), Paragraph(rl(cap), S["cap"])]))
    elif k == "refs":
        for ref in blk[1]:
            F.append(Paragraph(html.escape(ref, quote=False), S["ref"]))
    elif k == "pagebreak":
        F.append(PageBreak())
doc.build(F)
import re as _re

words = sum(len(_re.findall(r"[A-Za-z0-9]+", b[1])) for b in DOC if b[0] in ("p", "h1", "h2", "title")) + sum(
    len(_re.findall(r"[A-Za-z0-9]+", x)) for b in DOC if b[0] == "bullets" for x in b[1]
)
print("wrote", OUT + ".docx and .pdf; body words ~", words)
